from __future__ import annotations

import csv
import json
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

from export_service import _cell, _json_default, safe_export_filename

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
PDF_ROWS_PER_TABLE = 1000


def _safe(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=_json_default)
    return _cell(value)


def write_csv(path: Path, columns: list[str], rows: Iterable[dict[str, Any]], cancelled: Callable[[], bool]) -> int:
    processed = 0
    with Path(path).open("w", encoding="utf-8", newline="", buffering=1024 * 1024) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(columns)
        for row in rows:
            if cancelled():
                raise InterruptedError("Export cancelled.")
            writer.writerow([_cell(row.get(column)) if isinstance(row, dict) else "" for column in columns])
            processed += 1
    return processed


def write_json(path: Path, columns: list[str], rows: Iterable[dict[str, Any]], total_rows: int, cancelled: Callable[[], bool]) -> int:
    processed = 0
    first = True
    with Path(path).open("w", encoding="utf-8", buffering=1024 * 1024) as handle:
        handle.write('{"success":true,"columns":')
        handle.write(json.dumps(columns, ensure_ascii=False, default=_json_default, separators=(",", ":")))
        handle.write(',"total_rows":')
        handle.write(str(max(0, int(total_rows))))
        handle.write(',"rows":[')
        for row in rows:
            if cancelled():
                raise InterruptedError("Export cancelled.")
            if not first:
                handle.write(",")
            handle.write(json.dumps(row if isinstance(row, dict) else {}, ensure_ascii=False, default=_json_default, separators=(",", ":")))
            first = False
            processed += 1
        handle.write("]}")
    return processed


def write_xlsx(path: Path, columns: list[str], rows: Iterable[dict[str, Any]], cancelled: Callable[[], bool]) -> int:
    try:
        from openpyxl import Workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError("Excel export requires openpyxl.") from exc

    if len(columns) > EXCEL_MAX_COLUMNS:
        raise ValueError(f"Excel supports at most {EXCEL_MAX_COLUMNS:,} columns per worksheet.")

    wb = Workbook(write_only=True)
    try:
        sheet_number = 1
        ws = wb.create_sheet(f"Data_{sheet_number}")
        ws.append(columns)
        rows_in_sheet = 1
        processed = 0
        for row in rows:
            if cancelled():
                break
            if rows_in_sheet >= EXCEL_MAX_ROWS:
                sheet_number += 1
                ws = wb.create_sheet(f"Data_{sheet_number}")
                ws.append(columns)
                rows_in_sheet = 1
            ws.append([_safe(row.get(column)) if isinstance(row, dict) else "" for column in columns])
            rows_in_sheet += 1
            processed += 1
        wb.save(str(path))
        return processed
    finally:
        wb.close()


def write_pdf(path: Path, title: str, columns: list[str], rows: Iterable[dict[str, Any]], cancelled: Callable[[], bool]) -> int:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    except ModuleNotFoundError as exc:
        raise RuntimeError("PDF export requires reportlab.") from exc

    if not columns:
        raise ValueError("The report contains no columns to export.")

    doc = SimpleDocTemplate(
        str(path), pagesize=landscape(A4), rightMargin=8 * mm, leftMargin=8 * mm,
        topMargin=10 * mm, bottomMargin=10 * mm, title=title,
    )
    styles = getSampleStyleSheet()
    processed = 0
    iterator = iter(rows)

    def table_chunks() -> Iterator[Any]:
        nonlocal processed
        while True:
            if cancelled():
                raise InterruptedError("Export cancelled.")
            chunk: list[list[Any]] = [[Paragraph(str(column), styles["Heading5"]) for column in columns]]
            while len(chunk) < PDF_ROWS_PER_TABLE + 1:
                if cancelled():
                    raise InterruptedError("Export cancelled.")
                try:
                    row = next(iterator)
                except StopIteration:
                    break
                chunk.append([str(_safe(row.get(column))) if isinstance(row, dict) else "" for column in columns])
                processed += 1
            if len(chunk) <= 1:
                return
            table = LongTable(chunk, repeatRows=1, splitByRow=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f6")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d0d5dd")),
                ("FONTSIZE", (0, 0), (-1, -1), 6),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]))
            yield table
            if len(chunk) < PDF_ROWS_PER_TABLE + 1:
                return

    class _LazyFlowables(list):
        """Expose one PDF flowable at a time so completed chunks can be released."""
        def __init__(self) -> None:
            super().__init__([Paragraph(title or "Report", styles["Title"]), Spacer(1, 3 * mm)])
            self._chunks = iter(table_chunks())
            self._done = False

        def _ensure(self) -> None:
            if super().__len__() or self._done:
                return
            try:
                super().append(next(self._chunks))
            except StopIteration:
                self._done = True

        def __len__(self):
            self._ensure()
            return super().__len__()

        def __getitem__(self, index):
            self._ensure()
            return super().__getitem__(index)

        def __delitem__(self, index):
            super().__delitem__(index)
            self._ensure()

    doc.build(_LazyFlowables())
    return processed


def write_package(
    path: Path,
    title: str,
    columns: list[str],
    row_factory: Callable[[], Iterable[dict[str, Any]]],
    total_rows: int,
    cancelled: Callable[[], bool],
) -> int:
    """Build a ZIP package from independently streamed export artifacts.

    Each artifact gets a fresh ResultStore iterator. No complete result is
    materialized in memory. The package itself is written directly to disk.
    """
    with tempfile.TemporaryDirectory(prefix="report_export_") as temp_dir:
        root = Path(temp_dir)
        csv_path = root / "report.csv"
        json_path = root / "report.json"
        xlsx_path = root / "report.xlsx"
        pdf_path = root / "report.pdf"
        write_csv(csv_path, columns, row_factory(), cancelled)
        write_json(json_path, columns, row_factory(), total_rows, cancelled)
        write_xlsx(xlsx_path, columns, row_factory(), cancelled)
        write_pdf(pdf_path, title, columns, row_factory(), cancelled)
        if cancelled():
            raise InterruptedError("Export cancelled.")
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
            for artifact in (csv_path, json_path, xlsx_path, pdf_path):
                archive.write(artifact, artifact.name)
    return max(0, int(total_rows))
